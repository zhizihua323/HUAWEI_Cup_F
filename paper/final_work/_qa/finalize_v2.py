"""Read-only editorial validation then publish the already rendered PDF.
AI assisted: OpenAI Codex / GPT-6 Astra / OpenAI / 2026-09-03.
No model fitting, raw data access or scientific-result mutation.
"""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
import pdfplumber,json,re,hashlib,csv,shutil,sys
qa=Path(__file__).resolve().parent;out=qa.parent;root=out.parents[1]
folder=qa/sys.argv[1]
docpath=out/'FINAL_MANUSCRIPT_V2.docx';pdfpath=folder/'FINAL_MANUSCRIPT_V2.pdf'
doc=Document(docpath)
with ZipFile(docpath) as z:
    xml=E.fromstring(z.read('word/document.xml'));st=E.fromstring(z.read('word/styles.xml'))
ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
doc_text='\n'.join(xml.xpath('//w:t/text()',namespaces=ns))
with pdfplumber.open(pdfpath) as d:
    pages=[p.extract_text() or '' for p in d.pages]
    pagewidths=[(float(p.width),float(p.height)) for p in d.pages]
alltext='\n'.join(pages)
src=(out/'FINAL_MANUSCRIPT_V2_INTEGRATED.md').read_text(encoding='utf-8')
build=json.loads((qa/'v2_build_audit.json').read_text(encoding='utf-8'))
checks=[]
def check(name,passed,detail):checks.append({'check':name,'status':'PASS' if passed else 'FAIL','detail':detail})
check('A4',all(abs(w-595)<2 and abs(h-842)<2 for w,h in pagewidths),pagewidths[0])
check('abstract_two_pages','针对问题三' in pages[1] and '目' in pages[2] and '录' in pages[2],{'abstract_pages':[1,2],'toc_page':3})
check('continuous_footer',all(p.rstrip().splitlines()[-1]==str(i) for i,p in enumerate(pages,1)),len(pages))
check('no_placeholders',not any(s in doc_text for s in ['〔REF-','〔FIG-','@EQ','@TABLE','待接入','G1_PENDING']),True)
check('anonymous_body',not any(s in doc_text for s in ['28762','参赛学校','队员姓名','参赛队号']),True)
check('no_header',not xml.xpath('//w:headerReference',namespaces=ns),True)
check('all_display_math_editable',len(xml.xpath('//m:oMath',namespaces=ns))==27,27)
check('no_empty_sum_operand',all(len(n)>0 for n in xml.xpath('//m:nary/m:e',namespaces=ns)),True)
check('no_rendered_box',not any(c in alltext for c in ['□','�']),True)
check('tables_and_figures',len(doc.tables)==9 and len(doc.inline_shapes)==5,{'tables':len(doc.tables),'figures':len(doc.inline_shapes)})
check('reference_count',build['references']==18,build['citation_order'])
check('five_keywords',len(re.search(r'@KEYWORDS (.*)',src).group(1).split('；'))==5,5)
check('science_qualifiers',all(s in src for s in ['操作性质量代理','人工结果没有显示稳定正向一致性','SCENARIO_ONLY_UNVALIDATED','B2、B4和B5','比例分母也为零','质量不识别且不优化']),True)
anchors=['272505','261086','261067','0.032696','0.3544081081063713','1.6897975629820348','0.3539803206065571','1.2403055835426349','0.339976581941082','0.2798781285468448','0.5695341857475174','0.078248576','1.993850677','3.553996156','0.625922700','24.925757775','2.609142026','5.202388053','299.893000000','2.143211279','−3.971072','−1.933709','−0.488178']
check('frozen_anchor_text',all(a in src for a in anchors),len(anchors))
# TOC entries must match current rendered heading page map, not a prior build.
actual=json.loads((qa/'v2_heading_pages.json').read_text(encoding='utf-8'))
bad=[]
for p in doc.paragraphs:
    if p.style.name in ('TOC 1','TOC 2'):
        name,page=p.text.rsplit('\t',1)
        if page!=str(actual.get(name)):bad.append([name,page,actual.get(name)])
check('toc_exact_page_numbers',not bad,bad)
styles={s.style_id:s for s in doc.styles}
for name,size,east in [('Title',16,'黑体'),('Heading 1',14,'黑体'),('Heading 2',12,'宋体'),('Normal',12,'宋体'),('Caption',12,'宋体')]:
    s=doc.styles[name];f=s._element.xpath('./w:rPr/w:rFonts')[0]
    check('style_'+name,s.font.size.pt==size and f.get('{'+ns['w']+'}eastAsia')==east,{'size':s.font.size.pt,'eastAsia':f.get('{'+ns['w']+'}eastAsia')})
paths=[docpath,pdfpath,out/'FINAL_MANUSCRIPT_V2_SOURCE.md',out/'FINAL_MANUSCRIPT_V2_INTEGRATED.md']
paths += [Path(f['source']) for f in build['figures']]
paths += [root/'paper'/x for x in ['FINAL_SCIENTIFIC_FREEZE.md','FINAL_PAPER_PATCH_CONTRACT.md','FINAL_RESULT_INDEX_V2.md']]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
manifest=[{'path':str(p.relative_to(root)),'sha256':sha(p),'bytes':p.stat().st_size} for p in paths]
report={'status':'PASS' if all(c['status']=='PASS' for c in checks) else 'FAIL','pages':len(pages),'checks':checks,'science_executed':False,'scope':'Editorial artifact QA, not a new scientific acceptance run','rendering':'LibreOffice at D:/LibreOffice/program/soffice.exe, via render_docx.py','remaining_author_action':'Read and approve manuscript; reconcile historical source-file AI model attribution before submission packaging.'}
(qa/'V2_FINAL_CHECKS.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(qa/'V2_OUTPUT_MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
if report['status']!='PASS':
    print(json.dumps([c for c in checks if c['status']=='FAIL'],ensure_ascii=False));sys.exit(2)
shutil.copy2(pdfpath,out/'FINAL_MANUSCRIPT_V2.pdf')
print(json.dumps({'status':report['status'],'pages':len(pages),'checks':len(checks),'pdf':str(out/'FINAL_MANUSCRIPT_V2.pdf')},ensure_ascii=False))
