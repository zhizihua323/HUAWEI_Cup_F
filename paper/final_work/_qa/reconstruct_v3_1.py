"""Editorial DOCX build only. Does not execute any scientific pipeline.
AI assisted: OpenAI Codex, GPT-6 Astra, OpenAI, released 2026-09-03.
Inputs are the frozen-results manuscript and independent presentation assets.
"""
from pathlib import Path
import re, json, csv, copy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as E

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'paper/final_work'
QA=OUT/'_qa'
SRC=OUT/'FINAL_MANUSCRIPT_V3_1_SOURCE.md'
NS='http://www.w3.org/1998/Math/MathML'
def elem(tag,*children,text=None,**attrs):
    el=E.Element('{'+NS+'}'+tag,**attrs)
    if text is not None: el.text=text
    for c in children: el.append(c)
    return el
symbols={'alpha':'α','beta':'β','eta':'η','lambda':'λ','rho':'ρ','tau':'τ','phi':'φ','theta':'θ','Theta':'Θ','Delta':'Δ','varepsilon':'ε','partial':'∂','in':'∈','leq':'≤','geq':'≥','ldots':'…','cdot':'·','times':'×','mathsf':'','min':'min','max':'max','log':'log','inf':'inf','sum':'∑'}
class MathParser:
    def __init__(self,s): self.s=s.replace('\\left.','').replace('\\right.','');self.i=0
    def ws(self):
        while self.i<len(self.s) and self.s[self.i].isspace():self.i+=1
    def group(self):
        self.ws()
        if self.i<len(self.s) and self.s[self.i]=='{':
            self.i+=1;return self.seq('}')
        return self.atom()
    def atom(self):
        self.ws()
        if self.i>=len(self.s):return elem('mrow')
        c=self.s[self.i];self.i+=1
        if c=='{':return self.seq('}')
        if c=='\\':
            m=re.match(r'[A-Za-z]+',self.s[self.i:])
            if not m:
                ch=self.s[self.i];self.i+=1
                if ch in ',!; ':return elem('mspace',width='0.12em')
                return elem('mo',text=('‖' if ch=='|' else ch))
            name=m.group();self.i+=len(name)
            if name in ('left','right'):return self.atom()
            if name=='frac':return elem('mfrac',self.group(),self.group())
            if name=='sqrt':return elem('msqrt',self.group())
            if name in ('bar','overline','widehat','hat'):
                return elem('mover',self.group(),elem('mo',text='¯' if name in ('bar','overline') else '^'),accent='true')
            if name in ('mathrm','mathsf','mathcal','mathbb'):
                el=self.group()
                for n in el.iter():
                    if n.tag.endswith('mi'): n.set('mathvariant','normal')
                return el
            if name in ('quad','qquad'):return elem('mspace',width='0.7em' if name=='quad' else '1em')
            if name in symbols:
                el=elem('mo' if name in ('in','leq','geq','ldots','partial','sum','cdot','times') else 'mi',text=symbols[name])
                if name in ('min','max','log','inf'):el.set('mathvariant','normal')
                return el
            raise ValueError('unknown math command '+name)
        if c.isdigit():
            m=re.match(r'[\d.]*',self.s[self.i:]);v=c+m.group();self.i+=len(m.group());return elem('mn',text=v)
        return elem('mi' if c.isalpha() else 'mo',text=c)
    def seq(self,end=None):
        out=[]
        while self.i<len(self.s):
            self.ws()
            if self.i>=len(self.s):break
            if end and self.s[self.i]==end:self.i+=1;break
            a=self.atom();sub=sup=None;self.ws()
            while self.i<len(self.s) and self.s[self.i] in '_^':
                flag=self.s[self.i];self.i+=1
                if flag=='_':sub=self.group()
                else:sup=self.group()
                self.ws()
            if sub is not None and sup is not None:a=elem('msubsup',a,sub,sup)
            elif sub is not None:a=elem('msub',a,sub)
            elif sup is not None:a=elem('msup',a,sup)
            out.append(a)
        return elem('mrow',*out)

XSL=E.XSLT(E.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
def mathxml(tex):
    m=elem('math',MathParser(tex).seq(),display='block')
    result=XSL(m).getroot()
    # MML2OMML emits empty n-ary operands for a bare TeX summation.
    # Move the following summand into the operand, preserving expression order.
    for nary in result.findall('.//'+qn('m:nary')):
        operand=nary.find(qn('m:e'))
        if operand is not None and len(operand)==0:
            nxt=nary.getnext()
            while nxt is not None:
                t=''.join(nxt.itertext())
                if t.startswith((',', ';', '=')):break
                following=nxt.getnext();nxt.getparent().remove(nxt);operand.append(nxt);nxt=following
            assert len(operand)>0,tex
    # Ensure all math runs use the required readable body size.
    for mr in result.findall('.//'+qn('m:r')):
        wr=OxmlElement('w:rPr');fo=OxmlElement('w:rFonts');fo.set(qn('w:ascii'),'Cambria Math');fo.set(qn('w:hAnsi'),'Cambria Math');wr.append(fo)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'24');wr.append(sz)
        mr.insert(1 if mr.find(qn('m:rPr')) is not None else 0,wr)
    return result

text=SRC.read_text(encoding='utf-8').replace('四問','四问')
text=text.replace('10的18次方、20次方和22次方','10¹⁸、10²⁰和10²²')
text=text.replace('覆盖率99.9927%。A2和A3','覆盖率99.9927%。连续特征只用校准集的七域等权分位数确定归一化边界，留出与扩展数据不参与变换估计；其余11项规则和3项目标相关性指标保留为诊断，避免无方向依据的强制融合。A2和A3',1)
text=text.replace('未显示稳定正向一致性。因此','未显示稳定正向一致性。60条封存样本中46条完整、1条部分有效，其余记录无法无歧义恢复，统计按成对有效分母计算，不重抽样、不补评分。因此',1)
text=text.replace('但跨规模迁移仍有明显误差。','但跨规模迁移仍有明显误差。该配比模型在训练内部选择结构和正则化强度，在冻结后进行同尺度部署检验与跨尺度运输评价，两类证据分别报告。',1)
text=text.replace('该关系只在B6及其训练质量范围内成立。','该关系只在B6及其训练质量范围内成立。固定规模组内方向、相同测试键的交叉验证、200次重采样和参数识别检查共同支持这一限定性关系；B7重复点不增加证据，B8反向关系单列冲突。',1)
text=text.replace('最高预算的数据量达到支持上界。','最高预算的数据量达到支持上界。求解中保留参数与数据的经验范围，以活跃约束解释配置变化；上下文训练成本系数在长度30000处相等，但只称成本结构转移，不解释为能力相变。',1)
text=text.replace('不将损失下降换算为Benchmark提升。','不将损失下降换算为Benchmark提升。历史贡献分析与未来分位数情景使用不同资格的模型；同一未来时点更高算力对应更高情景分数，部分任务随时距增加下降来自未验证的负时间项，不能解释为能力必然衰退。',1)
text=text.replace('平台显示名称与可公开核验的发布版本应分开登记，不以猜测补造发布日期。具体工具信息及仍需核对的字段见随稿工具披露清单；未核实字段不作为已确认信息。','两模型官方发布日期分别为2026年7月9日和2026年9月3日，依据OpenAI官方更新记录核验[@openai_tools]。早期日志未逐一记录底层模型，故不为单个旧文件指定未经证据确认的模型归属。')
note='AI辅助说明：OpenAI Codex（GPT-5.6 Sol，OpenAI，2026年7月9日；GPT-6 Astra，OpenAI，2026年9月3日）用于代码草拟与结果整理；数值由冻结脚本和独立核验确定，使用范围见第9章[@openai_tools]。'
for heading in ['## 3.5 质量结果及其边界','## 4.5 分来源验证与结果解释','## 5.3 数值求解与主配置','## 6.3 历史关联模型及分解']:
    text=text.replace(heading,heading+'\n@NOTE '+note)
# Save the actual editorial source consumed by this build.
(OUT/'FINAL_MANUSCRIPT_V3_1_INTEGRATED.md').write_text(text,encoding='utf-8')

records={r['reference_key']:r for r in csv.DictReader((ROOT/'paper/final_references/reference_verification_final.csv').open(encoding='utf-8-sig'))}
records['openai_tools']={'authors_verified':'OpenAI','title':'GPT-5.6 Sol（2026-07-09）与GPT-6 Astra（2026-09-03）模型发布记录','year':'2026','venue':'OpenAI API Changelog','official_url':'https://developers.openai.com/api/docs/changelog','entry_type':'misc'}
citation_order=[]
for m in re.finditer(r'\[@([^\]]+)\]',text):
    for k in m.group(1).split(';'):
        assert k in records,k
        if k not in citation_order:citation_order.append(k)
def citations(s):
    return re.sub(r'\[@([^\]]+)\]',lambda m:'['+','.join(str(citation_order.index(k)+1) for k in m.group(1).split(';'))+']',s)

doc=Document(QA/'official_template_converted.docx')
body=doc._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)
sec=doc.sections[0]
sec.different_first_page_header_footer=False
sec.header_distance=Cm(.5);sec.footer_distance=Cm(.8)
for ref in list(sec._sectPr.findall(qn('w:headerReference'))):sec._sectPr.remove(ref)
for el in list(sec._sectPr.findall(qn('w:pgNumType'))):sec._sectPr.remove(el)
pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'1');sec._sectPr.append(pg)
# Remove template grid so body follows genuine single line spacing.
for el in list(sec._sectPr.findall(qn('w:docGrid'))):sec._sectPr.remove(el)
for footer in (sec.footer,sec.first_page_footer,sec.even_page_footer):
    for c in list(footer._element):footer._element.remove(c)
    p=footer.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)

def font(obj,size=12,bold=False,east='宋体'):
    obj.font.name='Times New Roman';obj.font.size=Pt(size);obj.font.bold=bold;obj.font.color.rgb=RGBColor(0,0,0)
    rp=obj._element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for key,val in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',east),('cs','Times New Roman')]:rf.set(qn('w:'+key),val)
def pstyle(name,size=12,bold=False,east='宋体'):
    if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st=doc.styles[name];font(st,size,bold,east)
    pf=st.paragraph_format;pf.line_spacing=1;pf.space_before=Pt(0);pf.space_after=Pt(3);pf.first_line_indent=Pt(24);pf.widow_control=True
    return st
pstyle('Normal')
pstyle('Title',16,True,'黑体')
for level in (1,2,3):
    st=pstyle('Heading '+str(level),14 if level==1 else 12,True,'黑体' if level==1 else '宋体')
    pf=st.paragraph_format;pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.space_before=Pt(12 if level==1 else 8);pf.space_after=Pt(6)
    pf.page_break_before=False
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
for n in ['Caption','TOC 1','TOC 2','Reference','Note']:
    st=pstyle(n);st.paragraph_format.first_line_indent=Pt(0)
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Caption'].paragraph_format.space_after=Pt(6)
doc.styles['Note'].paragraph_format.space_after=Pt(4)
doc.styles['Reference'].paragraph_format.left_indent=Pt(24)
doc.styles['Reference'].paragraph_format.first_line_indent=Pt(-24)
for n in ('TOC 1','TOC 2'):
    pf=doc.styles[n].paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=Pt(12);pf.keep_with_next=False;pf.keep_together=False;pf.widow_control=False;pf.left_indent=Pt(0 if n=='TOC 1' else 18)
    pf.tab_stops.add_tab_stop(Cm(16.45),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
font(doc.styles['TOC 1'],12,True)
# No template identity, author metadata, comments or running heading.
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='算力约束下大语言模型的数据质量评价与资源配置'
doc.core_properties.subject='数学建模论文提交候选版';doc.core_properties.comments=''
for el in list(doc.settings._element.findall(qn('w:updateFields'))):doc.settings._element.remove(el)
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings._element.append(upd)
for el in list(doc.settings._element.findall(qn('w:trackRevisions'))):doc.settings._element.remove(el)

def write_runs(p,s,style=None,bold=False):
    s=citations(s)
    # Superscript numeric citations; other text remains body font.
    parts=re.split(r'(\[\d+(?:,\d+)*\]|n_d,Q|(?<![A-Za-z0-9_])[A-Za-zŜρτλ̄]+(?:_[A-Za-z0-9τ]+)+)',s)
    for part in parts:
        if style!='Reference' and (re.fullmatch(r'[A-Za-zŜρτλ̄]+_[A-Za-z0-9τ]+',part) or part=='n_d,Q') and part not in ('Q_valid',):
            base,sub=part.split('_',1)
            r=p.add_run(base);font(r,12,bold);r.font.italic=True
            r=p.add_run(sub);font(r,12,bold);r.font.subscript=True
        else:
            r=p.add_run(part);font(r,12,bold)
            if style!='Reference' and re.fullmatch(r'\[\d+(?:,\d+)*\]',part):r.font.superscript=True
def para(s,style=None,align=None):
    p=doc.add_paragraph(style=style);p.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    write_runs(p,s,style)
    return p
def heading(s,level=1):
    p=doc.add_paragraph(s,style='Heading '+str(level));return p
def blank_break():
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)

table_widths={1:[2.5,3.1,10.9],2:[3.0,7.0,6.5],3:[3.2,11.4,1.9],4:[2.2,7.3,7.0],5:[4.1,4.0,8.4],6:[1.4,2.5,3.4,9.2],7:[3.0,4.4,4.6,4.5],8:[4.0,5.0,3.0,4.5],9:[1.5,7.0,8.0]}
def table(title,rows):
    p=para(title,'Caption',WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_with_next=True
    num=int(re.match(r'表(\d+)',title).group(1));widths=table_widths[num]
    tb=doc.add_table(rows=0,cols=len(rows[0]));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    for j,w in enumerate(widths):tb.columns[j].width=Cm(w)
    pr=tb._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    for i,values in enumerate(rows):
        cells=tb.add_row().cells
        trpr=tb.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,value in enumerate(values):
            cell=cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side,v in [('top','55'),('bottom','55'),('left','75'),('right','75')]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),v);el.set(qn('w:type'),'dxa');marg.append(el)
            cp.append(marg)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F3F5');cp.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
            p.paragraph_format.keep_with_next=i<len(rows)-1
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if len(value)>14 else WD_ALIGN_PARAGRAPH.CENTER
            write_runs(p,value,bold=i==0)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);font(p.add_run(''),3)

figmap={'fig01_framework.png':'fig01_v3_framework.png','fig02_q1_quality_validation.png':'fig02_v3_q1_quality.png','fig03_q1_regmix_transport.png':'fig03_v3_regmix_transport.png','fig04_q2_fit_validation.png':'fig04_v3_q2_validation.png','fig05_q4_scenarios.png':'fig05_v3_q4_scenarios.png'}
from PIL import Image
def figure(name,cap):
    path=ROOT/'paper/final_figures/submission_v3'/figmap[name]
    assert path.exists(),path
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
    w,h=Image.open(path).size
    width=16.5
    if h/w*width>21.0:width=21.0*w/h
    p.add_run().add_picture(str(path),width=Cm(width))
    p=para(cap,'Caption',WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_with_next=False
    fig_sources.append({'caption':cap,'source':str(path),'insert_width_cm':width,'height_cm':width*h/w})

def eq(tex,index):
    # Math paragraph with a right-aligned equation number; no equation tables.
    p=doc.add_paragraph();pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=Pt(4);pf.space_after=Pt(5);pf.keep_together=True
    pf.tab_stops.add_tab_stop(Cm(8.25),WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(16.45),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t');p._p.append(mathxml(tex));p.add_run('\t('+str(index)+')')

lines=text.splitlines();title=lines[0][2:]
p=doc.add_paragraph(title,style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(12)
abstract=True;toc_done=False;eqn=0;fig_sources=[];headings=[]
allheads=[(1 if l.startswith('# ') else 2,l.lstrip('# ')) for l in lines[1:] if l.startswith('# ') or l.startswith('## ')]
allheads.append((1,'参考文献'))
i=1
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line=='@ABSTRACT':
        p=para('摘  要',align=WD_ALIGN_PARAGRAPH.CENTER);font(p.runs[0],12,True);p.paragraph_format.first_line_indent=Pt(0);continue
    if line.startswith('@KEYWORDS '):
        p=para('关键词：'+line[10:]);p.paragraph_format.first_line_indent=Pt(0)
        blank_break();p=doc.add_paragraph('目  录',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(0)
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='右键单击并选择“更新域”以生成目录';r.append(t);fld.append(r);p._p.append(fld)
        blank_break();abstract=False;continue
    if line.startswith('## '):heading(line[3:],2);continue
    if line.startswith('# '):heading(line[2:]);continue
    if line.startswith('@NOTE '):para(line[6:],'Note');continue
    if line.startswith('@EQ '):eqn+=1;eq(line[4:],eqn);continue
    if line.startswith('@FIG '):
        a,b=line[5:].split('|',1);figure(a.strip(),b.strip());continue
    if line.startswith('@TABLE '):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            rows.append([v.strip() for v in lines[i].strip().strip('|').split('|')]);i+=1
        table(line[7:],rows);continue
    if line=='@REFERENCES':
        heading('参考文献')
        for n,key in enumerate(citation_order,1):
            r=records[key];auth=r['authors_verified'].split('; ');a=', '.join(auth[:3])+(', et al' if len(auth)>3 else '')
            title=r['title'];venue=r.get('venue','');year=r.get('year','')
            if r.get('entry_type')=='article':
                loc=r.get('volume','')+('('+r['issue']+')' if r.get('issue') else '')
                loc+=(':'+r['pages_or_article'] if r.get('pages_or_article') else '')
                bib=f'{a}. {title}[J]. {venue}, {year}, {loc}.'
            elif r.get('entry_type')=='inproceedings':
                bib=f'{a}. {title}[C]//{venue}. {year}'+(': '+r['pages_or_article'] if r.get('pages_or_article') else '')+'.'
            else:
                bib=f'{a}. {title}[EB/OL]. {year}[2026-09-26]. {r.get("official_url","")}.'
            if r.get('doi'):bib+=' DOI: '+r['doi']+'.'
            para('['+str(n)+'] '+bib,'Reference',WD_ALIGN_PARAGRAPH.LEFT)
        continue
    p=para(line)
    if abstract:
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.keep_together=True

doc.save(OUT/'FINAL_MANUSCRIPT_V3_1.docx')
(QA/'v3_1_build_audit.json').write_text(json.dumps({'equations':eqn,'references':len(citation_order),'citation_order':citation_order,'figures':fig_sources,'source_chars':len(text),'heading_count':len(allheads),'science_executed':False,'template':'official_template_converted.docx','toc':'automatic Word TOC field, levels 1-2, hyperlinks enabled','intentional_deviations':['Anonymous submission copy begins at abstract page 1','Remove identity cover and template artwork','Official all-other-Chinese 12pt SimSun overrides smaller captions/tables and Hei H2']},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(OUT/'FINAL_MANUSCRIPT_V3_1.docx'),'equations':eqn,'references':len(citation_order),'chars':len(text)},ensure_ascii=False))
