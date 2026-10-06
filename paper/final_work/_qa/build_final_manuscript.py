from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parents[3]
TEMPLATE = ROOT / "paper" / "final_work" / "_qa" / "official_template_converted.docx"
SOURCE = ROOT / "paper" / "final_work" / "FINAL_MANUSCRIPT_SOURCE.md"
OUTPUT = ROOT / "paper" / "final_work" / "FINAL_MANUSCRIPT_V1.docx"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
Q4_EQUATION_DIR = ROOT / "paper" / "final_work" / "_qa" / "equations"


NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


TABLE_CAPTIONS = [
    "主要符号与资格",
    "作者盲审与质量代理的相关结果",
    "M0_B1参数与支持框",
    "分来源验证结果",
    "算力预算下的冻结条件最优配置",
    "任务级模型选择与端点分解",
    "未验证的12个月和24个月Benchmark情景",
]


FIGURE_CAPTIONS = {
    "FIG-00": "四问数据流、冻结接口与资格分层技术路线",
    "FIG-Q1-01": "主质量覆盖、扩展冲突稳定性与作者盲审效应量",
    "FIG-Q1-02": "领域配比同尺度检验与跨尺度运输诊断",
    "FIG-Q2-01": "B1来源内拟合与分来源验证",
    "FIG-Q3-01": "三档预算下的冻结条件最优配置",
    "FIG-Q3-02": "冻结上下文长度下的活跃约束敏感性",
    "FIG-Q4-01": "六任务12个月和24个月三情景结果",
}


EQUATION_MATHML = {
    "Q1-01": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi>Q</mi><mrow><mi>i</mi><mi>g</mi></mrow></msub><mo>=</mo><mfrac><mrow><msub><mi>z</mi><mrow><mi>i</mi><mi>g</mi><mn>1</mn></mrow></msub><mo>+</mo><mo>⋯</mo><mo>+</mo><msub><mi>z</mi><mrow><mi>i</mi><mi>g</mi><msub><mi>m</mi><mi>g</mi></msub></mrow></msub></mrow><msub><mi>m</mi><mi>g</mi></msub></mfrac><mo>,</mo><mspace width="1em"/><msub><mi>Q</mi><mi>i</mi></msub><mo>=</mo><mfrac><mrow><msub><mi>Q</mi><mrow><mi>i</mi><mn>1</mn></mrow></msub><mo>+</mo><msub><mi>Q</mi><mrow><mi>i</mi><mn>2</mn></mrow></msub><mo>+</mo><msub><mi>Q</mi><mrow><mi>i</mi><mn>3</mn></mrow></msub></mrow><mn>3</mn></mfrac></mrow></math>""",
    "Q2-01": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi>L</mi><mn>0</mn></msub><mo>(</mo><msub><mi>N</mi><mi>B</mi></msub><mo>,</mo><msub><mi>D</mi><mi>B</mi></msub><mo>)</mo><mo>=</mo><mi>E</mi><mo>+</mo><mi>A</mi><msubsup><mi>N</mi><mi>B</mi><mrow><mo>−</mo><mi>α</mi></mrow></msubsup><mo>+</mo><mi>B</mi><msubsup><mi>D</mi><mi>B</mi><mrow><mo>−</mo><mi>β</mi></mrow></msubsup></mrow></math>""",
    "Q2-02": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi>Δ</mi><mi>Q</mi></msub><mo>(</mo><msub><mi>Q</mi><mi>B</mi></msub><mo>)</mo><mo>=</mo><mo>−</mo><msub><mi>k</mi><mi>add</mi></msub><mo>(</mo><msub><mi>Q</mi><mi>B</mi></msub><mo>−</mo><mn>0.6</mn><mo>)</mo></mrow></math>""",
    "Q2-03": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msubsup><mi>L</mi><mi>gen</mi><mi>s</mi></msubsup><mo>=</mo><msub><mi>L</mi><mn>0</mn></msub><mo>+</mo><msubsup><mi>ρ</mi><mi>Q</mi><mi>s</mi></msubsup><msub><mi>Δ</mi><mi>Q</mi></msub><mo>(</mo><msub><mi>h</mi><mi>s</mi></msub><mo>(</mo><msub><mi>Q</mi><mi>A</mi></msub><mo>)</mo><mo>)</mo><mo>+</mo><msubsup><mi>τ</mi><mi>p</mi><mi>s</mi></msubsup><msup><mi>c</mi><mi>T</mi></msup><mo>(</mo><mi>p</mi><mo>−</mo><msub><mi>p</mi><mn>0</mn></msub><mo>)</mo></mrow></math>""",
    "Q2-04": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><mfrac><mrow><mo>∂</mo><mi>L</mi></mrow><mrow><mo>∂</mo><msub><mi>N</mi><mi>B</mi></msub></mrow></mfrac><mo>=</mo><mo>−</mo><mi>α</mi><mi>A</mi><msubsup><mi>N</mi><mi>B</mi><mrow><mo>−</mo><mi>α</mi><mo>−</mo><mn>1</mn></mrow></msubsup><mo>,</mo><mspace width="1em"/><mfrac><mrow><mo>∂</mo><mi>L</mi></mrow><mrow><mo>∂</mo><msub><mi>D</mi><mi>B</mi></msub></mrow></mfrac><mo>=</mo><mo>−</mo><mi>β</mi><mi>B</mi><msubsup><mi>D</mi><mi>B</mi><mrow><mo>−</mo><mi>β</mi><mo>−</mo><mn>1</mn></mrow></msubsup></mrow></math>""",
    "Q3-01": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><mi>C</mi><mo>=</mo><mn>6</mn><mi>N</mi><mi>D</mi><mo>+</mo><mi>η</mi><mi>N</mi><mi>D</mi><mi>H</mi><mo>,</mo><mspace width="1em"/><mi>N</mi><mo>=</mo><msup><mn>10</mn><mn>9</mn></msup><msub><mi>N</mi><mi>B</mi></msub><mo>,</mo><mspace width="1em"/><mi>D</mi><mo>=</mo><msup><mn>10</mn><mn>9</mn></msup><msub><mi>D</mi><mi>B</mi></msub></mrow></math>""",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="BFBFBF", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, east="宋体", latin="Times New Roman", size=12, bold=None, italic=None, color="000000") -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)
    rfonts.set(qn("w:cs"), latin)


def format_body_paragraph(p, indent=True) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24) if indent else Pt(0)
    pf.widow_control = True
    for run in p.runs:
        set_run_font(run)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    npf = normal.paragraph_format
    npf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    npf.space_before = Pt(0)
    npf.space_after = Pt(0)

    for name, east, size, align, before, after in [
        ("Heading 1", "黑体", 14, WD_ALIGN_PARAGRAPH.CENTER, 12, 6),
        ("Heading 2", "黑体", 12, WD_ALIGN_PARAGRAPH.LEFT, 8, 3),
        ("Heading 3", "楷体", 12, WD_ALIGN_PARAGRAPH.LEFT, 6, 2),
    ]:
        st = styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st.paragraph_format.alignment = align
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.keep_together = True

    if "Figure Placeholder" not in styles:
        fp = styles.add_style("Figure Placeholder", WD_STYLE_TYPE.PARAGRAPH)
    else:
        fp = styles["Figure Placeholder"]
    fp.font.name = "Times New Roman"
    fp.font.size = Pt(10.5)
    fp.font.italic = True
    fp.font.color.rgb = RGBColor(89, 89, 89)
    fp._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    fp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(6)
    fp.paragraph_format.space_after = Pt(6)
    fp.paragraph_format.keep_together = True

    cap = styles["Caption"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(10.5)
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(3)
    cap.paragraph_format.keep_with_next = True


def remove_paragraph(paragraph) -> None:
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def extract_frontmatter(md: str):
    title = re.search(r"^#\s+(.+)$", md, re.M).group(1).strip()
    abstract_match = re.search(r"^## 摘要\s*\n(.*?)(?=\n\*\*关键词：\*\*)", md, re.S | re.M)
    abstract = re.sub(r"\n+", "\n", abstract_match.group(1).strip())
    keywords = re.search(r"\*\*关键词：\*\*\s*(.+)", md).group(1).strip()
    body = re.split(r"(?m)^# 1 引言\s*$", md, maxsplit=1)[1]
    body = "# 1 引言\n" + body
    return title, abstract, keywords, body


def clean_inline(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    text = text.replace("–", "-").replace("—", "-")
    return text.strip()


def add_paragraph_text(doc, text: str, style=None, indent=True, bold_lead=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and "：" in text:
        lead, rest = text.split("：", 1)
        r1 = p.add_run(lead + "：")
        set_run_font(r1, bold=True)
        r2 = p.add_run(rest)
        set_run_font(r2)
    else:
        r = p.add_run(clean_inline(text))
        set_run_font(r)
    if style is None:
        format_body_paragraph(p, indent=indent)
    return p


def add_toc(doc: Document) -> None:
    h = doc.add_paragraph("目录", style="Heading 1")
    h.paragraph_format.page_break_before = False
    entries = [
        "1 引言", "2 问题重述与总体思路", "3 数据 假设与符号",
        "4 问题一 数据质量评价与领域配比", "5 问题二 条件广义标度律",
        "6 问题三 算力预算下的条件优化", "7 问题四 Benchmark分解与未来情景",
        "8 统一验证与灵敏度", "9 模型评价", "10 改进与推广", "11 结论",
        "12 参考文献", "13 附录与复现说明",
    ]
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(entry)
        set_run_font(r, size=11)
    doc.add_page_break()


def mathml_to_omml(mathml: str):
    xslt = etree.XSLT(etree.parse(str(MML2OMML)))
    result = xslt(etree.fromstring(mathml.encode("utf-8")))
    return parse_xml(etree.tostring(result, encoding="unicode"))


def add_equation(doc: Document, eq_id: str, number: int) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(0.45), Inches(5.75), Inches(0.55)]
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 10, 20, 10, 20)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 2 else WD_ALIGN_PARAGRAPH.RIGHT
    middle = table.rows[0].cells[1].paragraphs[0]
    middle._p.append(mathml_to_omml(EQUATION_MATHML[eq_id]))
    nr = table.rows[0].cells[2].paragraphs[0].add_run(f"({number})")
    set_run_font(nr, size=10.5)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def add_equation_marker(doc: Document, marker: str, number: int) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(0.45), Inches(5.75), Inches(0.55)]
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 10, 20, 10, 20)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 2 else WD_ALIGN_PARAGRAPH.RIGHT
    mr = table.rows[0].cells[1].paragraphs[0].add_run(marker)
    set_run_font(mr, size=11)
    nr = table.rows[0].cells[2].paragraphs[0].add_run(f"({number})")
    set_run_font(nr, size=10.5)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def add_equation_image(doc: Document, image_name: str, number: int, width_in: float) -> None:
    image_path = Q4_EQUATION_DIR / image_name
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(0.45), Inches(5.75), Inches(0.55)]
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 10, 20, 10, 20)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 2 else WD_ALIGN_PARAGRAPH.RIGHT
    table.rows[0].cells[1].paragraphs[0].add_run().add_picture(str(image_path), width=Inches(width_in))
    nr = table.rows[0].cells[2].paragraphs[0].add_run(f"({number})")
    set_run_font(nr, size=10.5)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def add_equation_typographic(doc: Document, segments: list[tuple[str, bool]], number: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    for text, subscript in segments:
        r = p.add_run(text)
        set_run_font(r, east="Cambria Math", latin="Cambria Math", size=11)
        r.font.subscript = subscript
    nr = p.add_run(f"    ({number})")
    set_run_font(nr, size=10.5)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [clean_inline(c) for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows


def add_markdown_table(doc: Document, rows: list[list[str]], number: int) -> None:
    caption = doc.add_paragraph(style="Caption")
    r = caption.add_run(f"表 {number}  {TABLE_CAPTIONS[number - 1]}")
    set_run_font(r, size=10.5)
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, "D9E2F3")
            elif i % 2 == 0:
                set_cell_shading(cell, "F7F9FC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            text = row[j] if j < len(row) else ""
            run = p.add_run(text)
            set_run_font(run, size=9, bold=(i == 0))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)


def add_figure_slot(doc: Document, line: str, figure_number: int) -> None:
    match = re.search(r"〔(FIG-[^ ]+) 接入位〕", line)
    slot = match.group(1) if match else f"FIG-{figure_number:02d}"
    desc = FIGURE_CAPTIONS.get(slot, clean_inline(line))
    p = doc.add_paragraph(style="Figure Placeholder")
    r = p.add_run(f"图 {figure_number}  {desc}  [{slot} 待接入]")
    set_run_font(r, size=10.5, italic=True, color="595959")


def build_body(doc: Document, body: str) -> None:
    lines = body.splitlines()
    i = 0
    table_no = 0
    fig_no = 0
    eq_no = 0
    first_h1 = True
    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            table_no += 1
            add_markdown_table(doc, parse_table(block), table_no)
            continue
        m = re.fullmatch(r"\[EQ:([A-Z0-9-]+)\]", line)
        if m:
            eq_id = m.group(1)
            if eq_id == "Q4-01":
                formulas = (
                    (("M", False), ("0", True), (":  y = a", False)),
                    (("M", False), ("1", True), (":  y = a + b", False), ("C", True), (" log", False), ("10", True), ("(C)", False)),
                    (("M", False), ("2", True), (":  y = a + b", False), ("C", True), (" log", False), ("10", True), ("(C) + b", False), ("T", True), (" t", False)),
                )
                for segments in formulas:
                    eq_no += 1
                    add_equation_typographic(doc, list(segments), eq_no)
            else:
                eq_no += 1
                add_equation(doc, eq_id, eq_no)
            i += 1
            continue
        if line.startswith("〔FIG-"):
            fig_no += 1
            add_figure_slot(doc, line, fig_no)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(clean_inline(line[4:]), style="Heading 3")
            for run in p.runs:
                set_run_font(run, east="楷体", size=12, bold=True)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(clean_inline(line[3:]), style="Heading 2")
            for run in p.runs:
                set_run_font(run, east="黑体", size=12, bold=True)
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(clean_inline(line[2:]), style="Heading 1")
            if not first_h1:
                p.paragraph_format.page_break_before = True
            first_h1 = False
            for run in p.runs:
                set_run_font(run, east="黑体", size=14, bold=True)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph()
            item_no = re.match(r"^(\d+)\.\s+", line).group(1)
            r = p.add_run(item_no + ". " + clean_inline(re.sub(r"^\d+\.\s+", "", line)))
            set_run_font(r)
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-24)
            p.paragraph_format.space_after = Pt(0)
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph()
            r = p.add_run("• " + clean_inline(line[2:]))
            set_run_font(r)
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-12)
            p.paragraph_format.space_after = Pt(0)
            i += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("#", "|", "[EQ:", "〔FIG-", "- ")) or re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        add_paragraph_text(doc, "".join(paragraph_lines), indent=True)


def disable_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is not None:
        settings.remove(node)


def patch_frontmatter(doc: Document, title: str, abstract: str, keywords: str) -> None:
    paras = doc.paragraphs
    for idx in range(len(paras) - 1, 23, -1):
        remove_paragraph(paras[idx])

    paras = doc.paragraphs
    clear_paragraph(paras[17])
    paras[17].alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = paras[17].add_run("题  目：" + title)
    set_run_font(r, east="黑体", size=15, bold=True)
    paras[17].paragraph_format.line_spacing = 1.5

    clear_paragraph(paras[19])
    paras[19].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paras[19].add_run("摘  要")
    set_run_font(r, east="黑体", size=15, bold=True)

    abstract_para = paras[20]
    keyword_para = paras[23]
    remove_paragraph(paras[22])
    remove_paragraph(paras[21])
    clear_paragraph(abstract_para)
    r = abstract_para.add_run("".join(abstract.splitlines()))
    set_run_font(r, size=10.5)
    abstract_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract_para.paragraph_format.first_line_indent = Pt(21)
    abstract_para.paragraph_format.line_spacing = 1.15
    abstract_para.paragraph_format.space_after = Pt(0)

    clear_paragraph(keyword_para)
    r1 = keyword_para.add_run("关键词：")
    set_run_font(r1, east="黑体", size=11, bold=True)
    r2 = keyword_para.add_run(keywords)
    set_run_font(r2, size=11)
    keyword_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    keyword_para.paragraph_format.space_before = Pt(5)
    keyword_para.add_run().add_break(WD_BREAK.PAGE)


def main() -> int:
    if not TEMPLATE.exists() or not SOURCE.exists():
        raise FileNotFoundError("Template or manuscript source is missing")
    if not MML2OMML.exists():
        raise FileNotFoundError(f"Word MathML transform missing: {MML2OMML}")
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(str(OUTPUT))
    configure_styles(doc)
    title, abstract, keywords, body = extract_frontmatter(SOURCE.read_text(encoding="utf-8"))
    patch_frontmatter(doc, title, abstract, keywords)
    add_toc(doc)
    build_body(doc, body)
    disable_update_fields(doc)
    doc.core_properties.title = title
    doc.core_properties.subject = "华为杯中国研究生数学建模竞赛候选论文V1"
    doc.core_properties.author = ""
    doc.core_properties.keywords = keywords
    doc.core_properties.comments = "Built from the official Word template using frozen scientific interfaces."
    doc.save(str(OUTPUT))
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
