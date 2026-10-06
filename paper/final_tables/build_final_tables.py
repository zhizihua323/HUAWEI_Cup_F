"""Build nine frozen-result paper tables as CSV and DOCX fragments.

The script only selects and formats fields already present in frozen interfaces.
It performs no refit, threshold change, clipping, or scientific recomputation.
Run with the bundled document Python resolved by the Codex workspace runtime.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "paper" / "final_tables"
OUT.mkdir(parents=True, exist_ok=True)


def read_csv(rel: str) -> list[dict[str, str]]:
    with (ROOT / rel).open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def fmt(x, digits=6):
    if x in (None, "", "—"):
        return "—"
    try:
        val = float(x)
    except (TypeError, ValueError):
        return str(x)
    if abs(val) >= 1e5 or (0 < abs(val) < 1e-4):
        return f"{val:.6g}"
    return f"{val:.{digits}f}".rstrip("0").rstrip(".")


def write_csv(stem: str, headers: list[str], rows: list[list]) -> None:
    with (OUT / f"{stem}.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=8.5, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def write_docx(stem: str, number: str, title: str, headers: list[str], rows: list[list],
               note: str, landscape: bool = False) -> None:
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Cm(1.35); sec.bottom_margin = Cm(1.35)
    sec.left_margin = Cm(1.35); sec.right_margin = Cm(1.35)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.5)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{number} {title}")
    set_run_font(r, size=12, bold=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = True
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, h in enumerate(headers):
        cell = hdr.cells[j]; set_cell_shading(cell, "D9E5F0"); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        par = cell.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(str(h)); set_run_font(run, size=8.2, bold=True)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2:
                set_cell_shading(cells[j], "F5F7F8")
            par = cells[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(str(value)); set_run_font(run, size=7.7 if landscape else 8.1)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    run = p.add_run("注：" + note); set_run_font(run, size=7.4)
    doc.save(OUT / f"{stem}.docx")


def write_meta(stem: str, caption: str, sources: list[dict]) -> None:
    (OUT / f"{stem}_caption.md").write_text(caption.strip() + "\n", encoding="utf-8")
    payload = {"asset": stem, "transformation": "field_selection_and_formatting_only", "sources": sources}
    (OUT / f"{stem}_source_map.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def table01():
    stem = "table01_overview"
    headers = ["问题", "冻结主接口", "核心输出", "最终资格", "不可越过的解释边界"]
    rows = [
        ["Q1", "Q1_GAP + Q01C + T06", "质量代理；配比预测与运输", "CLOSED_WITH_LIMITATION", "Q 不是人工真值；跨尺度仅 TRANSPORT_ONLY"],
        ["Q2", "Q2_GAP + T06", "B1 来源内标度律；分来源验证", "CLOSED_WITH_LIMITATION", "不得称统一四变量模型已多源验证"],
        ["Q3", "T07", "10^18/10^20/10^22 FLOPs 条件最优 N、D", "CLOSED_WITH_LIMITATION", "Q 未识别；p 固定；10^22 的 D 达上界"],
        ["Q4", "Q4_GAP + T05/T08", "12/24 月 slow/baseline/upper 情景", "CLOSED_WITH_SCENARIO_LIMITATION", "全部 SCENARIO_ONLY_UNVALIDATED；Loss 转换为 0 次"],
    ]
    write_csv(stem, headers, rows)
    write_docx(stem, "表 1", "四问结果接口与最终资格", headers, rows,
               "本表只汇总冻结接口。资格优先级为 FINAL_SCIENTIFIC_FREEZE > FINAL_PAPER_PATCH_CONTRACT > FINAL_RESULT_INDEX_V2。", landscape=True)
    write_meta(stem, "表 1 | 四问结果接口与最终资格。", [{"file": "paper/FINAL_SCIENTIFIC_FREEZE.md", "fields": "Q1-Q4 status and qualification", "unit": "none", "denominator": "not applicable"}, {"file": "paper/FINAL_RESULT_INDEX_V2.md", "fields": "final main interfaces", "unit": "none", "denominator": "not applicable"}])


def table02():
    stem = "table02_q1_quality_definition"
    headers = ["项目", "冻结定义或数值", "单位/分母", "资格"]
    rows = [
        ["可用性组", "5 个特征，组内等权", "0–1", "严格完整案例"],
        ["知识组", "3 个特征，组内等权", "0–1", "严格完整案例"],
        ["教育推理组", "3 个特征，组内等权", "0–1", "严格完整案例"],
        ["Q_baseline", "三组等权；缺失不重分配", "0–1", "操作性质量代理"],
        ["物理记录", "272505", "条", "全量口径"],
        ["唯一键", "261086", "个", "去重口径"],
        ["主 Q 有效/缺失", "261067 / 19", "唯一键", "11 特征严格完整案例"],
    ]
    write_csv(stem, headers, rows)
    write_docx(stem, "表 2", "Q1 质量代理构造与主分母", headers, rows,
               "Q_baseline 只可称为基于冻结信号构造的操作性质量代理；不得称为语义质量金标准或人工真值。")
    write_meta(stem, "表 2 | Q1 质量代理构造与主分母。", [{"file": "solution/outputs/quality_q01c/20260924T215718+08/run_config.json", "fields": "primary_q_rule; group_rule", "unit": "0-1", "denominator": "11-feature strict complete case"}, {"file": "paper/FINAL_SCIENTIFIC_FREEZE.md", "fields": "physical rows; unique keys; Q valid/missing", "unit": "count", "denominator": "physical and unique reported separately"}])


def table03():
    stem = "table03_q1_extension_manual"
    stats = read_csv("diagnostics/TASK-G1/20260925T223500+08/manual_validation_statistics.csv")
    spearman = {r["analysis"]: r for r in stats if r["analysis"].startswith("spearman_Q_vs_")}
    headers = ["模块/指标", "有效分母 n", "估计值", "95% bootstrap 区间", "资格/说明"]
    rows = [
        ["A2 全局判据", "4 项", "1/4 通过", "—", "PARTIALLY_STABLE 的一部分"],
        ["A3 全局判据", "4 项", "3/4 通过", "—", "PARTIALLY_STABLE 的一部分"],
        ["A2/A3 域级扩展", "共同有效域均为 1", "—", "—", "NOT_EVALUABLE"],
    ]
    mapping = [
        ("spearman_Q_vs_overall_quality", "Q vs overall quality", "描述性人工核验"),
        ("spearman_Q_vs_readability", "Q vs readability", "方向不稳定"),
        ("spearman_Q_vs_completeness", "Q vs completeness", "描述性人工核验"),
        ("spearman_Q_vs_contamination", "Q vs contamination", "污染度越低越好"),
    ]
    for key, label, qual in mapping:
        r = spearman[key]
        rows.append([label, r["n_a"], fmt(r["effect_size"]), f"[{fmt(r['ci_low_2_5pct'])}, {fmt(r['ci_high_97_5pct'])}]", qual])
    rows.append(["作者评分解析", "60 seal", "46 完整 + 1 部分 + 10 无效 + 3 缺失", "—", "逐指标成对有效；不插补"])
    write_csv(stem, headers, rows)
    write_docx(stem, "表 3", "Q1 扩展稳定性与作者盲审", headers, rows,
               "相关系数为 Spearman ρ，区间基于 2000 次 bootstrap。人工结果未显示稳定正向一致性；不得用 AI 或规则补齐 13 条不可用记录。", landscape=True)
    write_meta(stem, "表 3 | Q1 扩展稳定性与作者盲审。", [{"file": "diagnostics/TASK-G1/20260925T204300+08/stability_decision.json", "fields": "A2/A3 criteria and domain evaluability", "unit": "criteria count", "denominator": "4 global criteria; common valid domains"}, {"file": "diagnostics/TASK-G1/20260925T223500+08/manual_validation_statistics.csv", "fields": "effect_size; ci bounds; n_a", "unit": "Spearman rho", "denominator": "pairwise complete author ratings"}, {"file": "diagnostics/TASK-G1/20260925T223500+08/manual_review_parse_audit.csv", "fields": "parse_status", "unit": "count", "denominator": "60 sealed review IDs"}])


def table04():
    stem = "table04_q1_regmix_transport"
    data = read_csv("solution/outputs/mixture/metrics.csv")
    wanted = {"test_1m", "test_60m", "test_1B"}
    rows = []
    for r in data:
        if r["model"] == "linear" and r["set"] in wanted:
            scale = {"test_1m": "1M", "test_60m": "60M", "test_1B": "1B"}[r["set"]]
            rows.append([scale, r["n"], fmt(r["rmse_all_domains"]), fmt(r["mae_all_domains"]), fmt(r["spearman_equal_domain_mean"]), r["role"], "SAME_SCALE_TEST" if scale == "1M" else "TRANSPORT_ONLY"])
    headers = ["尺度", "混合配比 n", "全域 RMSE", "全域 MAE", "等权域 Spearman", "冻结角色", "论文资格"]
    write_csv(stem, headers, rows)
    write_docx(stem, "表 4", "Q1 RegMix 同尺度检验与跨尺度运输", headers, rows,
               "误差单位为 Loss。60M 与 1B 沿用冻结线性关系，属于零样本跨尺度运输，不是跨尺度重新拟合或验证。", landscape=True)
    write_meta(stem, "表 4 | Q1 RegMix 同尺度检验与跨尺度运输。", [{"file": "solution/outputs/mixture/metrics.csv", "fields": "n; rmse_all_domains; mae_all_domains; spearman_equal_domain_mean; role", "unit": "Loss; rank correlation", "denominator": "frozen scale-specific test sets"}])


def table05():
    stem = "table05_q2_parameters_support"
    params = read_csv("diagnostics/TASK-T06E-INTEGRATE/20260925T110904+08/t07_parameter_table.csv")
    rows = []
    for r in params:
        if r["module"] == "M0_B1" and r["parameter"] in {"E", "A", "B", "alpha", "beta"}:
            rows.append([r["parameter"], fmt(r["value"], 15), f"[{fmt(r['p2_5'], 15)}, {fmt(r['p97_5'], 15)}]", r["source"], r["qualification"]])
    rows.extend([
        ["N_B 支持", "[0.070542, 11.965825]", "—", "B1", "十亿参数"],
        ["D_B 支持", "[0.134, 299.893]", "—", "B1", "十亿 token"],
    ])
    headers = ["参数/支持量", "冻结估计或范围", "2.5%–97.5%", "来源", "资格/单位"]
    write_csv(stem, headers, rows)
    write_docx(stem, "表 5", "Q2 M0 B1 参数与经验支持框", headers, rows,
               "模型为 L0=E+A·N_B^(-alpha)+B·D_B^(-beta)。所有参数逐位沿用冻结值，未重新拟合；区间为冻结接口中的来源条件区间。", landscape=True)
    write_meta(stem, "表 5 | Q2 M0 B1 参数与经验支持框。", [{"file": "diagnostics/TASK-T06E-INTEGRATE/20260925T110904+08/t07_parameter_table.csv", "fields": "M0_B1 parameter; value; p2_5; p97_5; qualification", "unit": "Loss; billion parameters; billion tokens", "denominator": "B1 source"}, {"file": "diagnostics/TASK-T07/20260925T113744+08/support_bounds.csv", "fields": "N_B and D_B lower/upper", "unit": "billion parameters; billion tokens", "denominator": "B1 empirical support"}])


def table06():
    stem = "table06_q2_validation"
    audit = {r["source_id"]: r for r in read_csv("diagnostics/TASK-G2/20260925T195336+08/support_oos_audit.csv")}
    metrics = read_csv("diagnostics/TASK-G2/20260925T195336+08/validation_metrics_by_source.csv")
    in_metrics = {(r["source_id"], r["subset"]): r for r in metrics}
    large = {r["source_id"]: r for r in read_csv("diagnostics/TASK-G2/20260925T195336+08/large_model_gt10b_audit.csv")}
    rows = []
    for sid in ["B2", "B3", "B4", "B5"]:
        a = audit[sid]; m = in_metrics[(sid, "IN_SUPPORT")]
        rows.append([sid, a["raw_rows"], f"{a['in_B1_box_rows']} / {int(a['finite_positive_N_D_rows'])-int(a['in_B1_box_rows'])}", fmt(m["rmse"]), m["final_qualification"], "支持内 RMSE；不跨来源合并"])
    rows.append(["B9", large["B9"]["raw_records"], "0 / 132", "不可计算真实误差", large["B9"]["final_qualification"], "无观测 Loss"])
    rows.append(["B10", large["B10"]["raw_records"], "0 / 128", fmt(large["B10"]["estimated_reference_rmse"]), large["B10"]["final_qualification"], "仅对估算参考"])
    headers = ["来源", "原始记录", "支持内 / OOS", "RMSE", "最终资格", "误差口径"]
    write_csv(stem, headers, rows)
    write_docx(stem, "表 6", "Q2 分来源验证结果与资格", headers, rows,
               "B3 为同源 Pythia 插值支持，非独立外部复制。B9/B10 均完全 OOS；B10 的 RMSE 只针对估算参考，不能作为真实误差。", landscape=True)
    write_meta(stem, "表 6 | Q2 分来源验证结果与资格。", [{"file": "diagnostics/TASK-G2/20260925T195336+08/support_oos_audit.csv", "fields": "raw_rows; in_B1_box_rows; finite_positive_N_D_rows", "unit": "count", "denominator": "source-specific records"}, {"file": "diagnostics/TASK-G2/20260925T195336+08/validation_metrics_by_source.csv", "fields": "IN_SUPPORT rmse; final_qualification", "unit": "Loss", "denominator": "source-specific support-in rows"}, {"file": "diagnostics/TASK-G2/20260925T195336+08/large_model_gt10b_audit.csv", "fields": "B9/B10 counts, estimated_reference_rmse, qualification", "unit": "count; Loss", "denominator": "all B9/B10 records"}])


def table07():
    stem = "table07_q3_budget_optima"
    data = read_csv("diagnostics/TASK-T07/20260925T113744+08/budget_scenario_optima.csv")
    rows = []
    for r in data:
        if r["scenario_id"] == "S00_NULL_M0_B1" and r["H"] == "2048":
            rows.append([f"{float(r['budget_flops']):.0e}", fmt(r["N_B"], 9), fmt(r["D_B"], 9), fmt(r["predicted_loss"], 9), r["active_constraints"], r["support_status"], r["Q_status"], r["p_status"]])
    headers = ["预算 FLOPs", "N_B", "D_B", "Loss", "活跃约束", "支持状态", "Q 状态", "p 状态"]
    write_csv(stem, headers, rows)
    write_docx(stem, "表 7", "Q3 三档预算正式主结果", headers, rows,
               "N_B、D_B 的单位分别为十亿参数和十亿 token。正式主场景为 S00_NULL_M0_B1、H=2048；1e22 的 D 达到支持上界。", landscape=True)
    write_meta(stem, "表 7 | Q3 三档预算正式主结果。", [{"file": "diagnostics/TASK-T07/20260925T113744+08/budget_scenario_optima.csv", "fields": "budget_flops; N_B; D_B; predicted_loss; active_constraints; support_status; Q_status; p_status", "unit": "FLOPs; billion parameters; billion tokens; Loss", "denominator": "scenario S00_NULL_M0_B1 and H=2048"}])


def table08():
    stem = "table08_q4_models_contributions"
    coef = {r["benchmark_task"]: r for r in read_csv("diagnostics/TASK-G4/20260925T195716+0800/taskwise_coefficients.csv")}
    contrib = {r["benchmark_task"]: r for r in read_csv("diagnostics/TASK-G4/20260925T195716+0800/scale_non_scale_contributions.csv")}
    tasks = ["IFEval", "BBH", "MATH Lvl 5", "GPQA", "MUSR", "MMLU-PRO", "six_task_mean_complete_only"]
    rows = []
    for t in tasks:
        c = coef[t]; z = contrib[t]
        rows.append(["六任务均值（辅助）" if t.startswith("six_") else t, c["selected_contribution_model"], "—" if c["selected_contribution_model"] == "M0_CONSTANT" else fmt(c["mean_bC_log10_compute"]), fmt(z["delta_scale_signed_points"]), fmt(z["delta_non_scale_signed_points"]), z["scale_association_abs_share"], "SCENARIO_ONLY_UNVALIDATED（未来）"])
    headers = ["任务", "选中贡献模型", "b_C", "Delta_scale", "Delta_non_scale", "比例", "未来数值资格"]
    write_csv(stem, headers, rows)
    write_docx(stem, "表 8", "Q4 选中贡献模型与端点分解", headers, rows,
               "Delta 单位为 Benchmark 点。两个正负 365 天窗口包含同一组 compute 有效模型，因此两项均为 0、比例 NOT_DEFINED；不得解释为规模没有作用。所有 M2_SCALE_TIME 均未通过升级门槛。", landscape=True)
    write_meta(stem, "表 8 | Q4 选中贡献模型与端点分解。", [{"file": "diagnostics/TASK-G4/20260925T195716+0800/taskwise_coefficients.csv", "fields": "selected_contribution_model; mean_bC_log10_compute", "unit": "Benchmark points per log10 compute", "denominator": "taskwise frozen fit"}, {"file": "diagnostics/TASK-G4/20260925T195716+0800/scale_non_scale_contributions.csv", "fields": "delta_scale_signed_points; delta_non_scale_signed_points; shares", "unit": "Benchmark points", "denominator": "frozen endpoint windows"}])


def table09():
    stem = "table09_q4_scenarios"
    data = read_csv("diagnostics/TASK-G4/20260925T195716+0800/forecast_12m_24m_taskwise.csv")
    data += read_csv("diagnostics/TASK-G4/20260925T195716+0800/forecast_auxiliary_composite.csv")
    tasks = ["IFEval", "BBH", "MATH Lvl 5", "GPQA", "MUSR", "MMLU-PRO", "six_task_mean_complete_only"]
    rows = []
    for t in tasks:
        vals = {}
        for r in data:
            if r["benchmark_task"] == t:
                vals[(int(r["horizon_months"]), r["scenario_id"])] = fmt(r["prediction_points"])
        def cell(h):
            return " / ".join(vals[(h, s)] for s in ["SLOW", "BASELINE", "UPPER_SENSITIVITY"])
        model = next(r["selected_contribution_model"] for r in data if r["benchmark_task"] == t)
        rows.append(["六任务均值（辅助）" if t.startswith("six_") else t, model, cell(12), cell(24), "SCENARIO_ONLY_UNVALIDATED"])
    headers = ["任务", "选中贡献模型", "12 月 SLOW / BASELINE / UPPER", "24 月 SLOW / BASELINE / UPPER", "资格"]
    write_csv(stem, headers, rows)
    write_docx(stem, "表 9", "Q4 十二和二十四个月三情景", headers, rows,
               "预测原点 2025-01-28；目标日期为 2026-01-28 和 2027-01-28；单位为 Benchmark 点。数值全部未截断，MATH Lvl 5 的 24 月负值原样保留。增长率来自数学网格，不是历史估计。", landscape=True)
    write_meta(stem, "表 9 | Q4 十二和二十四个月三情景。", [{"file": "diagnostics/TASK-G4/20260925T195716+0800/forecast_12m_24m_taskwise.csv", "fields": "benchmark_task; horizons; scenario_id; prediction_points; eligibility", "unit": "Benchmark points", "denominator": "6 primary tasks x 3 scenarios x 2 horizons"}, {"file": "diagnostics/TASK-G4/20260925T195716+0800/forecast_auxiliary_composite.csv", "fields": "auxiliary composite scenarios", "unit": "Benchmark points", "denominator": "1 auxiliary composite x 3 scenarios x 2 horizons"}])


def main():
    for fn in [table01, table02, table03, table04, table05, table06, table07, table08, table09]:
        fn()


if __name__ == "__main__":
    main()
