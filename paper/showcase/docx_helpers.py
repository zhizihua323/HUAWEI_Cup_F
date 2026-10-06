from __future__ import annotations
from pathlib import Path
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

def set_run_font(run, cn='SimSun', latin='Times New Roman', size=12, bold=None, italic=None, color=None):
    run.font.name=latin
    run._element.rPr.rFonts.set(qn('w:eastAsia'),cn)
    run._element.rPr.rFonts.set(qn('w:ascii'),latin)
    run._element.rPr.rFonts.set(qn('w:hAnsi'),latin)
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color: run.font.color.rgb=RGBColor(*color)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_margins(cell,top=40,start=60,bottom=40,end=60):
    tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for tag,val in (('top',top),('start',start),('bottom',bottom),('end',end)):
        node=OxmlElement(f'w:{tag}'); node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa'); mar.append(node)
    tcPr.append(mar)

def set_cell_text(cell,text,bold=False,size=9.0,center=False):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(str(text)); set_run_font(r,size=size,bold=bold); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def borders(table):
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),'888888'); borders.append(el)
    table._tbl.tblPr.append(borders)

def add_page_number_footer(section):
    section.footer.is_linked_to_previous=False; p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin'); b=OxmlElement('w:instrText'); b.set(qn('xml:space'),'preserve'); b.text=' PAGE '; c=OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'),'end')
    run._r.append(a); run._r.append(b); run._r.append(c); set_run_font(run,size=10.5)

def page_start(section,start=1):
    sectPr=section._sectPr; e=sectPr.find(qn('w:pgNumType'))
    if e is None: e=OxmlElement('w:pgNumType'); sectPr.append(e)
    e.set(qn('w:start'),str(start))

def configure_styles(doc):
    normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'SimSun'); normal.font.size=Pt(12); normal.paragraph_format.line_spacing=1.0; normal.paragraph_format.space_after=Pt(2)
    vals=[('Title','SimHei',18,True,WD_ALIGN_PARAGRAPH.CENTER,0,12),('Heading 1','SimHei',14,True,WD_ALIGN_PARAGRAPH.CENTER,10,6),('Heading 2','SimHei',12,True,WD_ALIGN_PARAGRAPH.LEFT,7,4),('Heading 3','SimHei',12,True,WD_ALIGN_PARAGRAPH.LEFT,5,3)]
    for name,cn,size,bold,align,before,after in vals:
        st=doc.styles[name]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'),cn); st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.alignment=align; st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def add_title(doc,text,size=16):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(text); set_run_font(r,cn='SimHei',size=size,bold=True)

def add_para(doc,text,first=True):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(2)
    if first: p.paragraph_format.first_line_indent=Pt(24)
    r=p.add_run(text); set_run_font(r,size=12); return p

def add_h1(doc,text): doc.add_heading(text,level=1)
def add_h2(doc,text): doc.add_heading(text,level=2)
def add_h3(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True
    r=p.add_run(text); set_run_font(r,cn='SimHei',size=12,bold=True)

def add_formula(doc,formula,num):
    p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.left_indent=Cm(1.0); p.paragraph_format.right_indent=Cm(0.2); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(3); p.paragraph_format.tab_stops.add_tab_stop(Cm(15.0),WD_ALIGN_PARAGRAPH.RIGHT)
    r=p.add_run(formula+'\t'); set_run_font(r,cn='Cambria Math',latin='Cambria Math',size=11.5); r2=p.add_run(f'({num})'); set_run_font(r2,size=11.5)

def add_figure(doc,path,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1); p.add_run().add_picture(str(path),width=Cm(15.1))
    cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(5); r=cap.add_run(caption); set_run_font(r,size=10.2)

def add_table(doc,t):
    cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_before=Pt(4); cap.paragraph_format.space_after=Pt(2); r=cap.add_run(t['title']); set_run_font(r,cn='SimHei',size=10.3,bold=True)
    table=doc.add_table(rows=1,cols=len(t['headers'])); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=True
    hdr=table.rows[0]
    for i,h in enumerate(t['headers']):
        set_cell_text(hdr.cells[i],h,bold=True,size=8.6,center=True); shade(hdr.cells[i],'E9EEF1')
    repeat_header(hdr)
    for row in t['rows']:
        cells=table.add_row().cells
        for i,val in enumerate(row): set_cell_text(cells[i],val,size=8.4,center=(i>0 and len(str(val))<=18))
    borders(table)
    note=doc.add_paragraph(); note.paragraph_format.first_line_indent=Pt(0); note.paragraph_format.space_before=Pt(1); note.paragraph_format.space_after=Pt(5); nr=note.add_run('注：'+t['note']); set_run_font(nr,size=9.0)

def add_toc(doc):
    p=doc.add_paragraph(); run=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin'); b=OxmlElement('w:instrText'); b.set(qn('xml:space'),'preserve'); b.text=' TOC \\o "1-2" \\h \\z \\u '; c=OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'),'separate'); d=OxmlElement('w:fldChar'); d.set(qn('w:fldCharType'),'end'); t=OxmlElement('w:t'); t.text='目录将在打开文档时更新'
    run._r.append(a); run._r.append(b); run._r.append(c); run._r.append(t); run._r.append(d)
