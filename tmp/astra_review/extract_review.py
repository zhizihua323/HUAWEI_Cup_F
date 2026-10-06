from pathlib import Path
from pypdf import PdfReader
from docx import Document
import json, hashlib
root=Path.cwd()
out=root/'tmp/astra_review'
out.mkdir(parents=True,exist_ok=True)
for name,path in [('v3',root/'paper/final_work/FINAL_MANUSCRIPT_V3.pdf'),('reference2020',root/'参考资料/2020优秀论文.pdf')]:
    reader=PdfReader(path)
    pages=[p.extract_text() or '' for p in reader.pages]
    (out/f'{name}.txt').write_text('\n\n'.join(f'=== PDF PAGE {i+1} ===\n{t}' for i,t in enumerate(pages)),encoding='utf-8')
    (out/f'{name}_pages.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
    print(name,'pages',len(pages),'sha256',hashlib.sha256(path.read_bytes()).hexdigest())
for path in [root/'paper/final_work/FINAL_MANUSCRIPT_V3.docx',*root.glob('附件2*.docx'),*root.glob('附件4*.docx')]:
    d=Document(path)
    texts=[p.text for p in d.paragraphs]
    for t in d.tables:
        texts.extend(' | '.join(c.text for c in r.cells) for r in t.rows)
    (out/(path.stem+'.txt')).write_text('\n'.join(texts),encoding='utf-8')
    print(path.name,'paragraphs',len(d.paragraphs),'tables',len(d.tables),'sha256',hashlib.sha256(path.read_bytes()).hexdigest())
