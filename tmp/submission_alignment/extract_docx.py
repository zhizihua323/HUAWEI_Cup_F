import json
import sys
from pathlib import Path

from docx import Document


def main() -> None:
    path = Path(sys.argv[1])
    mode = sys.argv[2] if len(sys.argv) > 2 else "text"
    doc = Document(path)
    if mode == "text":
        print("\n".join(p.text for p in doc.paragraphs))
        print("\n[TABLES]\n")
        for table in doc.tables:
            for row in table.rows:
                print(" | ".join(cell.text for cell in row.cells))
        return

    paragraphs = []
    headings = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        item = {"index": index, "style": style, "text": text}
        paragraphs.append(item)
        if style.startswith("Heading") or style.startswith("标题") or style == "Title":
            headings.append(item)
    tables = []
    for index, table in enumerate(doc.tables):
        tables.append({
            "index": index,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "first_row": " | ".join(c.text for c in table.rows[0].cells) if table.rows else "",
        })
    print(json.dumps({
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "headings": headings,
        "paragraphs": paragraphs,
        "tables": tables,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
